from __future__ import annotations

import json
import re
import tempfile
import unittest
from pathlib import Path

from shared.design_tokens import resolve_design_tokens, validate_custom_style
from shared.pptx_static_core import contrast_ratio

STYLES = (
    "Academic Rigorous",
    "Data Driven",
    "Modern Minimal",
    "Charcoal Editorial",
    "Midnight Business",
    "Ocean Tech",
    "Teal Trust",
    "Cherry Bold",
    "Creative Student",
    "Coral Energy",
    "Forest Moss",
    "Warm Terracotta",
)
ROOT = Path(__file__).resolve().parents[1]


class DesignTokenTests(unittest.TestCase):
    def test_catalog_has_three_categories_of_four_unique_styles(self) -> None:
        catalog = json.loads((ROOT / "references" / "design-tokens.json").read_text(encoding="utf-8"))
        self.assertEqual(12, len(catalog["styles"]))
        self.assertEqual(3, len(catalog["categories"]))
        flattened = [key for values in catalog["categories"].values() for key in values]
        self.assertEqual(12, len(flattened))
        self.assertEqual(12, len(set(flattened)))
        self.assertEqual(set(catalog["styles"]), set(flattened))
        self.assertTrue(all(len(values) == 4 for values in catalog["categories"].values()))

    def test_all_standard_styles_resolve_to_lightweight_safety_tokens(self) -> None:
        for style in STYLES:
            with self.subTest(style=style):
                tokens = resolve_design_tokens(style)
                for field in ("style_character", "palette", "backgrounds", "svg_reference"):
                    self.assertIn(field, tokens)
                for shared in ("geometry", "lines", "typography"):
                    self.assertIn(shared, tokens)
                self.assertNotIn("style_dna", tokens)
                self.assertEqual("16:9", tokens["geometry"]["slide_ratio"])

    def test_legacy_aliases_map_with_compatibility_warning(self) -> None:
        cases = {
            "Berry Cream": "Warm Terracotta",
            "berry-cream": "Warm Terracotta",
            "Sage Calm": "Forest Moss",
            "sage-calm": "Forest Moss",
        }
        for old, expected in cases.items():
            with self.subTest(old=old):
                tokens = resolve_design_tokens(old)
                self.assertEqual(expected, tokens["style_name"])
                self.assertTrue(tokens["compatibility_warnings"])

    def test_other_requires_and_resolves_complete_custom_reference(self) -> None:
        custom = {
            "style_character": "Quiet scientific field notes",
            "palette": {
                "canvas": "FFFFFF",
                "surface": "F5F5F5",
                "primary_text": "111111",
                "secondary_text": "555555",
                "primary_accent": "2563EB",
                "secondary_accent": "93C5FD",
            },
            "backgrounds": {
                "cover": "Blue field",
                "content": "White canvas",
                "section": "Pale blue field",
                "closing": "Blue field",
            },
            "svg_reference": {"name": "none", "usage": "No recurring SVG motif"},
        }
        self.assertEqual([], validate_custom_style(custom))
        tokens = resolve_design_tokens("Other", custom)
        self.assertEqual("Other", tokens["style_name"])
        self.assertEqual(custom["palette"], tokens["palette"])
        with self.assertRaises(ValueError):
            resolve_design_tokens("Other")
        low_contrast = json.loads(json.dumps(custom))
        low_contrast["palette"]["secondary_text"] = "EEEEEE"
        self.assertTrue(any("4.5:1" in item for item in validate_custom_style(low_contrast)))
        with_extra = json.loads(json.dumps(custom))
        with_extra["layout_preference"] = "two-column"
        self.assertTrue(any("unsupported fields" in item for item in validate_custom_style(with_extra)))

    def test_unknown_style_uses_safe_reference_and_warning(self) -> None:
        tokens = resolve_design_tokens("School template: green and gold")
        self.assertTrue(tokens["custom_style"])
        self.assertEqual("School template: green and gold", tokens["style_character"])
        self.assertTrue(tokens["compatibility_warnings"])
        self.assertEqual("16:9", tokens["geometry"]["slide_ratio"])

    def test_all_styles_pass_role_aware_contrast(self) -> None:
        for style in STYLES:
            with self.subTest(style=style):
                palette = resolve_design_tokens(style)["palette"]
                for role in ("primary_text", "secondary_text"):
                    for background in ("canvas", "surface"):
                        self.assertGreaterEqual(contrast_ratio(palette[role], palette[background]), 4.5)
                for background in ("canvas", "surface"):
                    self.assertGreaterEqual(contrast_ratio(palette["primary_accent"], palette[background]), 3.0)

    def test_every_style_has_a_contrast_safe_dark_companion(self) -> None:
        seen: set[tuple[str, ...]] = set()
        for style in STYLES:
            with self.subTest(style=style):
                dark = resolve_design_tokens(style)["dark_palette"]
                self.assertEqual(
                    {"canvas", "surface", "primary_text", "secondary_text",
                     "primary_accent", "secondary_accent"},
                    set(dark),
                )
                for role in ("primary_text", "secondary_text"):
                    for background in ("canvas", "surface"):
                        self.assertGreaterEqual(contrast_ratio(dark[role], dark[background]), 4.5)
                for background in ("canvas", "surface"):
                    self.assertGreaterEqual(contrast_ratio(dark["primary_accent"], dark[background]), 3.0)
                seen.add(tuple(sorted(dark.items())))
        # Twelve styles must not collapse onto identical dark schemes.
        self.assertEqual(len(STYLES), len(seen))

    def test_custom_style_dark_companion_is_optional_but_validated(self) -> None:
        custom = {
            "style_character": "Quiet scientific field notes",
            "palette": {
                "canvas": "FFFFFF", "surface": "F5F5F5", "primary_text": "111111",
                "secondary_text": "555555", "primary_accent": "2563EB", "secondary_accent": "93C5FD",
            },
            "backgrounds": {"cover": "c", "content": "c", "section": "s", "closing": "c"},
            "svg_reference": {"name": "none", "usage": "n"},
        }
        # Omitted: derived, not an error.
        self.assertEqual([], validate_custom_style(custom))
        derived = resolve_design_tokens("Other", custom)["dark_palette"]
        self.assertGreaterEqual(contrast_ratio(derived["primary_text"], derived["canvas"]), 4.5)
        # Supplied but unreadable: rejected.
        broken = json.loads(json.dumps(custom))
        broken["dark_palette"] = {
            "canvas": "0F172A", "surface": "1E293B", "primary_text": "22303F",
            "secondary_text": "2A3847", "primary_accent": "1D4ED8", "secondary_accent": "30447E",
        }
        errors = validate_custom_style(broken)
        self.assertTrue(any("4.5:1" in item for item in errors))
        good = json.loads(json.dumps(custom))
        good["dark_palette"] = {
            "canvas": "0F172A", "surface": "1E293B", "primary_text": "F1F5F9",
            "secondary_text": "CBD5E1", "primary_accent": "93C5FD", "secondary_accent": "30447E",
        }
        self.assertEqual([], validate_custom_style(good))

    def test_style_files_have_exact_four_fields_and_match_tokens(self) -> None:
        expected_fields = ["Style character", "Palette", "Background reference", "SVG reference"]
        style_dir = ROOT / "skills" / "sp-deck" / "references" / "visual-styles"
        self.assertEqual(12, len(list(style_dir.glob("*.md"))))
        for style in STYLES:
            with self.subTest(style=style):
                tokens = resolve_design_tokens(style)
                text = (style_dir / f"{tokens['style_key']}.md").read_text(encoding="utf-8")
                actual = re.findall(r"^- \*\*([^:]+):\*\*", text, flags=re.MULTILINE)
                self.assertEqual(expected_fields, actual)
                for value in tokens["palette"].values():
                    self.assertIn(str(value), text)
                self.assertIn(tokens["svg_reference"]["name"], text)
                # The cover/section/closing scheme must be stated in every style file.
                self.assertIn(tokens["dark_palette"]["canvas"], text)

    def test_every_style_has_a_distinct_font_pairing(self) -> None:
        cjk_title = {"Microsoft YaHei", "SimHei", "DengXian", "KaiTi"}
        cjk_body = {"Microsoft YaHei", "DengXian", "DengXian Light", "SimSun", "FangSong", "KaiTi"}
        pairings = set()
        for style in STYLES:
            with self.subTest(style=style):
                typography = resolve_design_tokens(style)["typography"]
                self.assertTrue(typography["title_font"])
                self.assertTrue(typography["body_font"])
                self.assertIn(typography["cjk_title_font"], cjk_title)
                self.assertIn(typography["cjk_body_font"], cjk_body)
                pairings.add(
                    (
                        typography["title_font"],
                        typography["body_font"],
                        typography["cjk_title_font"],
                        typography["cjk_body_font"],
                    )
                )
        # Font personality must differ across styles; no two share a full pairing.
        self.assertEqual(len(STYLES), len(pairings))

    def test_cjk_font_postprocessor_writes_ea_typeface(self) -> None:
        from pptx import Presentation
        from pptx.util import Inches

        from shared.pptx_runtime import apply_cjk_fonts

        with tempfile.TemporaryDirectory() as tmp:
            pptx_path = Path(tmp) / "tiny.pptx"
            deck = Presentation()
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
            for face in ("Cambria", "Calibri"):
                run = box.text_frame.paragraphs[0].add_run()
                run.text = "标题正文"
                run.font.name = face
            deck.save(pptx_path)

            written = apply_cjk_fonts(
                pptx_path, {"Cambria": "SimHei", "Calibri": "Microsoft YaHei"}
            )
            self.assertEqual(2, written)
            reloaded = Presentation(pptx_path)
            ea_faces = []
            for shape in reloaded.slides[0].shapes:
                for para in shape.text_frame.paragraphs:
                    for run in para.runs:
                        ea = run.font._rPr.find(
                            "{http://schemas.openxmlformats.org/drawingml/2006/main}ea"
                        )
                        if ea is not None:
                            ea_faces.append(ea.get("typeface"))
            self.assertIn("SimHei", ea_faces)
            self.assertIn("Microsoft YaHei", ea_faces)

    def test_every_standard_style_declares_a_visual_language(self) -> None:
        allowed = {
            "rule": {"bracket", "left-rail", "underline-left", "top-band", "slash"},
            "panel": {"outlined", "flush", "soft-fill", "edge-band"},
            "motif_at": {"corner-tr", "corner-bl", "edge-right"},
            "chart": {"columns", "line", "bars"},
            "decor": {"restrained", "balanced", "expressive"},
            "pattern": {"none", "dots", "waves", "grid"},
            "emphasis_marker": {"dot", "square", "slash", "chevron", "number", "dash", "leaf"},
            "cover_band": {"none", "bottom-band", "side-band", "corner-block"},
        }
        languages = set()
        for style in STYLES:
            with self.subTest(style=style):
                language = resolve_design_tokens(style)["visual_language"]
                self.assertEqual(
                    {
                        "rule", "panel", "radius", "motif_at", "chart", "decor",
                        "pattern", "emphasis_marker", "cover_band",
                    },
                    set(language),
                )
                for key, choices in allowed.items():
                    self.assertIn(language[key], choices)
                self.assertIsInstance(language["radius"], (int, float))
                self.assertTrue(0 <= language["radius"] <= 0.25)
                languages.add(json.dumps(language, sort_keys=True))
        # Languages must differentiate the styles: no two styles share one full language.
        self.assertEqual(len(STYLES), len(languages))


if __name__ == "__main__":
    unittest.main()
