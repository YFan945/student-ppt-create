"""Regression tests for audit findings in generation and visual QA."""
from __future__ import annotations

import hashlib
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from pptx import Presentation
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "skills" / "sp-deck" / "scripts"))

import builder_packet  # noqa: E402
import pptx_quality_gate_v071 as quality  # noqa: E402
import pptx_rendered_check as rendered  # noqa: E402
import run_gates  # noqa: E402
from pipeline.core import Stage, collect, qa_input_fingerprint  # noqa: E402

from shared.quality_tiers import tier_policy  # noqa: E402


class GateIntegrityTests(unittest.TestCase):
    def test_subprocess_crash_cannot_reuse_old_success(self):
        with tempfile.TemporaryDirectory() as tmp:
            work = Path(tmp)
            (work / "qa-rendered.json").write_text('{"ok":true,"issues":[]}', encoding="utf-8")
            args = run_gates.parse_args(["--pptx", str(work / "deck.pptx")])
            with patch.object(run_gates.subprocess, "run", return_value=subprocess.CompletedProcess([], 1, "", "crash")):
                report = run_gates.run(args, work)
            self.assertFalse(report["ok"])
            self.assertEqual(1, report["counts"]["critical"])

    def test_bad_reports_and_minor_only_failure_are_blockers(self):
        for payload, code in (("[]", 0), ("invalid JSON", 0), ('{"ok":false,"issues":[]}', 0), ('{"ok":false,"issues":[{"severity":"minor"}]}', 2), ('{"ok":true,"issues":[]}', 2)):
            with self.subTest(payload=payload), tempfile.TemporaryDirectory() as tmp:
                report_path = Path(tmp) / "report.json"
                def execute(_argv, report_path=report_path, payload=payload, code=code, **_kwargs):
                    report_path.write_text(payload, encoding="utf-8")
                    return subprocess.CompletedProcess([], code, "", "")
                gates, problems = {}, []
                with patch.object(run_gates.subprocess, "run", side_effect=execute):
                    run_gates._qa_gate("test", "test.py", [], gates, problems, report_path)
                self.assertFalse(gates["test"]["ok"])
                self.assertTrue(any(item["severity"] in run_gates.BLOCKING for item in problems))

    def test_critic_fix_survives_quality_pipeline_and_packet(self):
        with tempfile.TemporaryDirectory() as tmp:
            work = Path(tmp)
            pptx = work / "deck.pptx"
            pptx.write_bytes(b"fixture")
            finding = {"code": "overlap", "severity": "critical", "message": "chart overlap", "element": "chart1", "fix": "move chart to the right", "repair_level": "implementation"}
            review = work / "review.json"
            review.write_text(json.dumps({"pptx_sha256": quality.sha256_file(pptx), "slides": [{"slide": 1, "visual_structure": "chart", "scores": {field: 8 for field in quality.SCORE_FIELDS}, "issues": [finding]}]}), encoding="utf-8")
            result = quality.validate_visual_report(review, pptx, 1, policy=tier_policy("rigorous"))
            report_path = work / "quality.json"
            def execute(_argv):
                report_path.write_text(json.dumps(result), encoding="utf-8")
                return subprocess.CompletedProcess([], 2, "", "")
            with patch("pipeline.core._runner", side_effect=execute):
                _ok, problems, _binding = collect(Stage("quality", [], report_path, "quality"))
            projected = builder_packet.project_blocker(next(item for item in problems if item["code"] == "overlap"))
            for key in ("element", "fix", "repair_level"):
                self.assertEqual(finding[key], projected[key])

    def test_schema_requires_actionable_blockers(self):
        schema = json.loads((ROOT / "references" / "visual-review.schema.json").read_text(encoding="utf-8"))
        import jsonschema
        validator = jsonschema.Draft202012Validator(schema)
        report = {"pptx_sha256": "0"*64, "slides": [{"slide": 1, "visual_structure": "chart", "scores": {field: 8 for field in quality.SCORE_FIELDS}, "issues": [{"severity": "major", "code": "overlap", "message": "fix"}]}]}
        self.assertTrue(list(validator.iter_errors(report)))
        report["slides"][0]["issues"][0].update(element="chart1", fix="move chart", repair_level="implementation")
        self.assertFalse(list(validator.iter_errors(report)))

    def make_deck(self, path, sizes):
        deck = Presentation()
        for title_size, body_size in sizes:
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            for name, text, size, y in (("Title", "Title", title_size, 0.5), ("Body", "Body", body_size, 2)):
                shape = slide.shapes.add_textbox(Inches(0.5), Inches(y), Inches(8), Inches(0.5))
                shape.name = name
                shape.text = text
                shape.text_frame.paragraphs[0].runs[0].font.size = Pt(size)
        deck.save(path)

    def test_large_cover_cannot_hide_flat_body_page(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "deck.pptx"
            self.make_deck(path, [(44, 22), (22, 22)])
            report = rendered.check_pptx(path, {})
            self.assertFalse(report["ok"])
            self.assertEqual([2], [item["slide"] for item in report["issues"] if item["code"] == "title-body-ratio"])

    def test_intentional_whitespace_is_advisory(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "deck.pptx"
            self.make_deck(path, [(33, 22)])
            report = rendered.check_pptx(path, {})
            self.assertTrue(report["ok"], report)
            self.assertEqual("advisory", next(item["severity"] for item in report["issues"] if item["code"] == "dead-space"))

    def test_rich_title_runs_do_not_dominate_the_body_font(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "deck.pptx"
            self.make_deck(path, [(33, 22)])
            deck = Presentation(path)
            title, body = deck.slides[0].shapes
            title.name, body.name = "Text0", "Text1"
            title.text_frame.clear()
            for _ in range(10):
                run = title.text_frame.paragraphs[0].add_run()
                run.text, run.font.size = "Title", Pt(33)
            deck.save(path)
            report = rendered.check_pptx(path, {})
            self.assertTrue(report["ok"], report["issues"])
            self.assertEqual(1.5, report["metrics"]["page_title_body_ratio"]["slide-1"])

    def test_registry_report_must_bind_current_deck(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "deck.pptx"
            self.make_deck(path, [(33, 22)])
            sidecar = Path(str(path) + ".registry-report.json")
            data = {"pptx_sha256": "0"*64, "warnings": [{"slide": 1, "code": "content_dead_zone", "message": "gap"}]}
            sidecar.write_text(json.dumps(data), encoding="utf-8")
            self.assertEqual("registry-report-stale", rendered.registry_sidecar_findings(path)[0]["code"])
            data["pptx_sha256"] = hashlib.sha256(path.read_bytes()).hexdigest()
            sidecar.write_text(json.dumps(data), encoding="utf-8")
            self.assertEqual("registry-content_dead_zone", rendered.registry_sidecar_findings(path)[0]["code"])

    def test_registry_change_invalidates_qa_cache(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "deck.pptx"
            path.write_bytes(b"fixture")
            before = qa_input_fingerprint(path, None, None, [])
            sidecar = Path(str(path) + ".registry-report.json")
            sidecar.write_text("{}", encoding="utf-8")
            self.assertNotEqual(before, qa_input_fingerprint(path, None, None, []))
            sidecar.write_text("invalid JSON", encoding="utf-8")
            self.assertEqual("registry-report-invalid", rendered.registry_sidecar_findings(path)[0]["code"])

    def test_background_and_footer_do_not_hide_content_gap(self):
        from pptx.enum.shapes import MSO_SHAPE
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "deck.pptx"
            self.make_deck(path, [(33, 22)])
            original = rendered.check_pptx(path, {})["metrics"]["bottom_whitespace"]
            deck = Presentation(path)
            slide = deck.slides[0]
            slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, deck.slide_width, deck.slide_height)
            footer = slide.shapes.add_textbox(0, deck.slide_height-Inches(0.4), Inches(3), Inches(0.3))
            footer.name = "Footer"
            footer.text = "Page 1"
            footer.text_frame.paragraphs[0].runs[0].font.size = Pt(11)
            deck.save(path)
            self.assertEqual(original, rendered.check_pptx(path, {})["metrics"]["bottom_whitespace"])

    def test_master_inherited_font_sizes_are_measured(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "deck.pptx"
            self.make_deck(path, [(33, 22)])
            deck = Presentation(path)
            for shape in deck.slides[0].shapes:
                shape.text_frame.paragraphs[0].runs[0].font.size = None
            master = deck.slide_master._element
            master.find(".//p:titleStyle/a:lvl1pPr/a:defRPr", rendered.NS).set("sz", "3300")
            master.find(".//p:otherStyle/a:lvl1pPr/a:defRPr", rendered.NS).set("sz", "2200")
            deck.save(path)
            metrics = rendered.check_pptx(path, {})["metrics"]
            self.assertEqual(22, metrics["font_min_pt"])
            self.assertEqual(1.5, metrics["page_title_body_ratio"]["slide-1"])

    def test_group_coordinates_are_resolved_on_the_slide(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "deck.pptx"
            deck = Presentation()
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            group = slide.shapes.add_group_shape()
            for y, name, size in ((1, "Title", 33), (2, "Body", 22)):
                shape = group.shapes.add_textbox(Inches(1), Inches(y), Inches(5), Inches(0.5))
                shape.name, shape.text = name, name
                shape.text_frame.paragraphs[0].runs[0].font.size = Pt(size)
            group.top = Inches(3)
            group.height = Inches(3)
            deck.save(path)
            report = rendered.check_pptx(path, {})
            self.assertAlmostEqual(0.2, report["metrics"]["bottom_whitespace"]["slide-1"])
            self.assertEqual(1.5, report["metrics"]["page_title_body_ratio"]["slide-1"])

    def test_delivery_smoke_fixture_cannot_pass_as_critic(self):
        from test_helpers import load_module
        matrix = load_module(ROOT / "scripts" / "scenario_render_matrix.py")
        with tempfile.TemporaryDirectory() as tmp:
            path, review = Path(tmp) / "deck.pptx", Path(tmp) / "review.json"
            path.write_bytes(b"fixture")
            matrix.write_visual_review(path, review, 1)
            report = json.loads(review.read_text(encoding="utf-8"))
            self.assertNotIn("scores", report["slides"][0])
            self.assertTrue(quality.visual_review_schema_issues(report))
