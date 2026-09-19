"""Regression coverage for orchestration boundaries, using real OOXML/delivery IO."""
from __future__ import annotations

import io
import json
import shlex
import subprocess
from contextlib import redirect_stdout

from pptx import Presentation
from test_ppt_pipeline import FakeRunner, PipelineTestCase, pp


class PipelineIntegrityTests(PipelineTestCase):
    def prepared(self):
        files = self.write_inputs()
        self.plan(files)
        pp.main(["build", "--work-dir", str(self.work), "--entry", str(self.entry())])
        self.render_evidence(files)
        return files

    def test_next_command_reaches_real_delivery_with_notes_and_one_preview_per_slide(self):
        files = self.prepared()
        deck = Presentation()
        deck.slides.add_slide(deck.slide_layouts[6])
        deck.save(files["pptx"])
        manifest = self.manifest()
        manifest["build"]["pptx"] = pp.bind(files["pptx"])
        manifest["render"] = {}
        pp.save_manifest(self.work, manifest)

        def render_runner(argv):
            from PIL import Image
            page = self.work / "render/slide-1.png"
            image = Image.new("RGB", (640, 360), "white")
            image.paste("navy", (0, 0, 640, 80))
            image.save(page)
            return subprocess.CompletedProcess(argv, 0, stdout=json.dumps({"pages": [str(page)]}), stderr="")

        pp._core._runner = render_runner
        self.assertEqual(pp.main(["render", "--work-dir", str(self.work)]), 0)
        self.render_evidence(files)
        fake = FakeRunner(self.work)
        calls = []
        results = []

        def runner(argv):
            if any("pptx_delivery_check_v08.py" in arg for arg in argv):
                calls.append(argv)
                result = pp.run_command([*argv, "--json"])
                results.append(json.loads(result.stdout))
                return result
            return fake(argv)

        pp._core._runner = runner
        output = io.StringIO()
        with redirect_stdout(output):
            pp.main(["next", "--work-dir", str(self.work), "--json"])
        command = json.loads(output.getvalue())["next_command"]
        # Parse the emitted, fully quoted command exactly (no added QA flags).
        argv = [arg.strip('"') for arg in shlex.split(command, posix=False)]
        self.assertEqual(pp.main(argv[2:]), 2)  # unrelated fake reports must NOT pass real delivery
        self.assertEqual(len(calls), 1)
        report = results[0]
        self.assertNotIn("notes", report["missing_expected_files"])
        self.assertNotIn("preview", report["missing_expected_files"])
        self.assertEqual(report["delivery_report"]["preview_page_coverage"], "1/1")
        self.assertEqual(calls[0].count("--preview"), 1)
        self.assertIn(str(self.work / "speaker-notes.md"), calls[0])

    def test_changed_image_refuses_qa(self):
        files = self.prepared()
        (self.work / "render/slide-1.png").write_bytes(b"tampered")
        self.assertEqual(pp.main(["qa", "--work-dir", str(self.work), "--visual-review", str(files["visual_review"])]), 2)

    def test_wrong_work_id_cannot_authorize_or_mirror(self):
        self.prepared()
        state = json.loads(self.workflow_state.read_text())
        state["work_id"] = "another-deck"
        self.workflow_state.write_text(json.dumps(state))
        with self.assertRaises(pp.RefusedError):
            pp.validate_manifest_authorization(self.manifest())
        with self.assertRaises(pp.RefusedError):
            pp.mirror_workflow_state(self.manifest(), "complete")
        self.assertEqual(json.loads(self.workflow_state.read_text())["work_id"], "another-deck")

    def test_outside_work_dir_and_output_traversal_are_refused(self):
        self.assertEqual(pp.main(["next", "--work-dir", str(self.work.parent.parent)]), 2)
        files = self.write_inputs()
        self.plan(files)
        self.assertEqual(pp.main(["build", "--work-dir", str(self.work), "--entry", str(self.entry()), "--output-name", "../escaped.pptx"]), 2)
        self.assertFalse((self.work.parent / "escaped.pptx").exists())

    def test_real_ooxml_unpack_edit_pack_preserves_original(self):
        files = self.write_inputs()
        source = self.work.parent / "original.pptx"
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[5])
        slide.shapes.title.text = "Original"
        deck.save(source)
        original = source.read_bytes()
        spec = json.loads(files["spec"].read_text())
        spec.update(source_deck=str(source), edit_intent="review-fix", preserve=["theme"], change_summary_required=True)
        files["spec"].write_text(json.dumps(spec))
        fake = FakeRunner(self.work)

        def runner(argv):
            if str(pp.PPTX_TOOL) in argv and any(command in argv for command in ["unpack", "pack"]):
                return pp.run_command(argv)
            return fake(argv)

        self.plan(files, runner)
        self.assertEqual(self.manifest()["mode"], "edit_ooxml")
        self.assertFalse((self.work / "deck.js").exists())
        part = self.work / "ooxml/ppt/slides/slide1.xml"
        part.write_text(part.read_text(encoding="utf-8").replace("Original", "Revised"), encoding="utf-8")
        self.assertEqual(pp.main(["build", "--work-dir", str(self.work)]), 0)
        self.assertEqual(source.read_bytes(), original)
        self.assertEqual(Presentation(files["pptx"]).slides[0].shapes.title.text, "Revised")

    def test_research_scope_requires_actual_runtime_receipt(self):
        files = self.write_inputs()
        spec = json.loads(files["spec"].read_text())
        spec["research_scope"] = "D"
        files["spec"].write_text(json.dumps(spec))
        pp._core._runner = FakeRunner(self.work)
        self.assertEqual(pp.main(["plan", "--work-dir", str(self.work), "--slide-spec", str(files["spec"]), "--validation-report", str(files["spec_report"]), "--art-direction", str(files["art"])]), 2)

    def test_source_image_is_rebuild_and_requires_analysis(self):
        files = self.write_inputs()
        source = self.work / "source.png"
        source.write_bytes(b"source fixture")
        spec = json.loads(files["spec"].read_text())
        spec["source_deck"] = str(source)
        files["spec"].write_text(json.dumps(spec))
        args = ["plan", "--work-dir", str(self.work), "--slide-spec", str(files["spec"]), "--validation-report", str(files["spec_report"]), "--art-direction", str(files["art"])]
        pp._core._runner = FakeRunner(self.work)
        self.assertEqual(pp.main(args), 2)
        (self.work / "source-analysis.md").write_text("Rebuild the supplied source while preserving approved content.")
        self.assertEqual(pp.main(args), 0)
        self.assertEqual(self.manifest()["mode"], "rebuild_from_source")
        self.assertIn("source_analysis", self.manifest()["inputs"])

    def test_second_work_does_not_reset_first_authorization(self):
        files = self.write_inputs()
        self.plan(files)
        first_work, first_state = self.work, self.workflow_state
        original = first_state.read_bytes()
        self.work = first_work.with_name("work-02")
        self.work.mkdir()
        self.summary = self.work / "production-summary.md"
        self.summary.write_text("another approval")
        self.workflow_state = self.work / "workflow-state.json"
        self.confirm_intake()
        self.plan(self.write_inputs())
        self.assertEqual(first_state.read_bytes(), original)
        pp.validate_manifest_authorization(pp.load_manifest(first_work))

    def test_mutated_notes_prevent_complete(self):
        files = self.prepared()
        pp._core._runner = FakeRunner(self.work)
        self.assertEqual(pp.main(["qa", "--work-dir", str(self.work), "--visual-review", str(files["visual_review"])]), 0)
        (self.work / "speaker-notes.md").write_text("Changed after QA")
        self.assertEqual(pp.main(["complete", "--work-dir", str(self.work)]), 2)
